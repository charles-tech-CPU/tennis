import openpyxl

XLSX = r'C:\Users\Utilisateur\Desktop\EXCEL\TENNIS\2026\20260914TENNIS 2026.xlsx'

def cell(ws, r, c):
    return ws.cell(row=r, column=c).value

def decode_side(ws, start_row, name_col, tag_col, round_col_fn, half_rounds):
    """Decode one half (16 rows) : round0 names/tags + round-by-round winner/score."""
    names = []
    for i in range(16):
        r = start_row + i
        names.append({'row': r, 'name': cell(ws, r, name_col), 'tag': cell(ws, r, tag_col)})

    rounds = {}
    for rnd in range(1, half_rounds + 1):
        col = round_col_fn(rnd)
        group = 1 << rnd
        results = []
        m = 0
        while start_row + m * group <= start_row + 15:
            top = start_row + m * group
            w_row = top + (1 << (rnd - 1)) - 1
            s_row = top + (1 << (rnd - 1))
            results.append({
                'group_top': top,
                'winner_name': cell(ws, w_row, col),
                'score': cell(ws, s_row, col),
            })
            m += 1
        rounds[rnd] = results
    return names, rounds

def decode_main_draw(ws, header_row, X):
    """32-slot tournament (drawSize<=32), mirrored 2-half layout, base column X (=tag_col_left).

    Chaque moitie (gauche/droite) ne dispose que de 3 colonnes dediees (R32,R16,QF)
    avant d'entrer en collision avec l'autre moitie : son propre tour 4 (demi-finale
    du cote) est donc REEMBOYTE dans la colonne du tour 3 (QF), a la ligne du milieu
    du demi-tableau complet (16 lignes). La VRAIE finale (tour 5, qui oppose les deux
    demi-finalistes) occupe la colonne centrale X+5, isolee entre les deux moities,
    egalement a la ligne du milieu.
    """
    name = cell(ws, header_row, X + 1)
    pts = [cell(ws, header_row, X + 1 + r) for r in range(1, 6)]
    start_row = header_row + 1

    left_names, left_rounds = decode_side(ws, start_row, X + 1, X,
                                           lambda r: X + 1 + r if r <= 3 else X + 4, 4)
    right_names, right_rounds = decode_side(ws, start_row, X + 9, X + 10,
                                             lambda r: X + 9 - r if r <= 3 else X + 6, 4)

    final_col = X + 5
    w_row = start_row + (1 << 3) - 1
    s_row = start_row + (1 << 3)
    final = {'winner_name': cell(ws, w_row, final_col), 'score': cell(ws, s_row, final_col)}

    return {
        'name': name, 'points': pts, 'start_row': start_row,
        'left_names': left_names, 'left_rounds': left_rounds,
        'right_names': right_names, 'right_rounds': right_rounds,
        'final': final,
    }

def decode_qualifying(ws, header_row, X):
    """X = colonne du tag de qualif (name='QUALIF' est a X+1, points a X+2/X+3)."""
    name_hdr = cell(ws, header_row, X + 1)  # 'QUALIF'
    pts = [cell(ws, header_row, X + 2), cell(ws, header_row, X + 3)]
    start_row = header_row + 1
    name_col = X + 1
    tag_col = X

    # Taille du tirage de qualifs variable selon le tournoi : on lit tant que la
    # colonne "nom" (ou le tag) n'est pas vide, en s'arretant sur un multiple de 4
    # (structure en groupes independants de 4).
    size = 0
    r = start_row
    while cell(ws, r, name_col) is not None or cell(ws, r, tag_col) is not None:
        size += 1
        r += 1
    size = (size // 4) * 4  # arrondi au multiple de 4 inferieur, par securite

    names = []
    for i in range(size):
        r = start_row + i
        names.append({'row': r, 'name': cell(ws, r, name_col), 'tag': cell(ws, r, tag_col)})

    rounds = {}
    for rnd in range(1, 3):
        col = name_col + rnd
        group = 1 << rnd
        results = []
        m = 0
        while start_row + m * group <= start_row + size - 1:
            top = start_row + m * group
            w_row = top + (1 << (rnd - 1)) - 1
            s_row = top + (1 << (rnd - 1))
            results.append({'group_top': top, 'winner_name': cell(ws, w_row, col), 'score': cell(ws, s_row, col)})
            m += 1
        rounds[rnd] = results
    return {'points': pts, 'start_row': start_row, 'size': size, 'names': names, 'rounds': rounds}

def print_main_draw(d):
    print(f"=== {d['name']} === points={d['points']}")
    print("-- Left half round0 --")
    for n in d['left_names']:
        print(f"  row{n['row']}: tag={n['tag']!r:6} name={n['name']}")
    for rnd, results in d['left_rounds'].items():
        print(f"-- Left round {rnd} --")
        for res in results:
            print(f"  group@{res['group_top']}: winner={res['winner_name']!r:20} score={res['score']!r}")
    print("-- Right half round0 --")
    for n in d['right_names']:
        print(f"  row{n['row']}: tag={n['tag']!r:6} name={n['name']}")
    for rnd, results in d['right_rounds'].items():
        print(f"-- Right round {rnd} --")
        for res in results:
            print(f"  group@{res['group_top']}: winner={res['winner_name']!r:20} score={res['score']!r}")
    print(f"-- FINALE -- winner={d['final']['winner_name']!r} score={d['final']['score']!r}")

def print_qualifying(d):
    print(f"=== QUALIF (size={d['size']}) === points={d['points']}")
    for n in d['names']:
        print(f"  row{n['row']}: tag={n['tag']!r:6} name={n['name']}")
    for rnd, results in d['rounds'].items():
        print(f"-- Q round {rnd} --")
        for res in results:
            print(f"  group@{res['group_top']}: winner={res['winner_name']!r:20} score={res['score']!r}")

def find_header_rows(ws, col=2, max_row=200):
    """Aide au reperage : liste (row, texte) pour chaque cellule texte non vide de
    la colonne donnee (2=B ou 20=T selon le cote de la mise en page) - permet de
    retrouver les lignes d'en-tete de chaque tournoi de la feuille (voir README
    section 9 pour la methode complete)."""
    out = []
    for r in range(1, max_row):
        v = ws.cell(row=r, column=col).value
        if v and isinstance(v, str) and v.strip():
            out.append((r, v.strip()))
    return out


if __name__ == '__main__':
    import sys
    # Usage:
    #   decode_bracket.py find <sheet_name>              -> liste les en-tetes (col B et col T)
    #   decode_bracket.py show <sheet_name> <header_row> <X>  -> decode un tournoi precis
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    if len(sys.argv) >= 3 and sys.argv[1] == 'find':
        ws = wb[sys.argv[2]]
        print("-- colonne B --")
        for r, v in find_header_rows(ws, col=2):
            print(r, v)
        print("-- colonne T (20) --")
        for r, v in find_header_rows(ws, col=20):
            print(r, v)
    elif len(sys.argv) >= 5 and sys.argv[1] == 'show':
        ws = wb[sys.argv[2]]
        header_row, X = int(sys.argv[3]), int(sys.argv[4])
        d = decode_main_draw(ws, header_row, X)
        print_main_draw(d)
        q = decode_qualifying(ws, header_row, X + 12)
        print_qualifying(q)
    else:
        print(__doc__ or "Usage: decode_bracket.py find|show <sheet> [header_row] [X]")
